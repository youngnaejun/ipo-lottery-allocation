"""Write Tables 1-7, B1-B4 and IA1-IA3 to a Word file from the Stata (and SAS) result files.

Usage: build_word_tables.py <stata_out_dir> <table_text.json> <out.docx>

Inputs in <stata_out_dir>: results_table1.csv, results_table2.csv, table2_median_params.csv and
results_{returns,panels,revisions,appendix_b,ppml}.csv written by stata/do/00_master.do.
Number formats follow the paper: regression coefficients and t/z statistics with three significant
digits (t/z in parentheses), stars on coefficients (*** p<0.01, ** p<0.05, * p<0.10), N with thousands
separators, fit statistics with three decimals; Table 1 and Table 2 with four decimals (percentages two).
Titles, panel headings and notes come from <table_text.json> (workflow/analysis/table_text.json: neutral wording).
"""
import csv
import json
import math
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

out_dir, text_json, out_docx = Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3])
TEXT = json.loads(text_json.read_text(encoding="utf-8"))


def text_for(prefix):
    return next(t for t in TEXT if t["caption"].startswith(prefix))


def sig3(x):
    """Three significant digits in fixed notation, keeping trailing zeros (e.g. 0.000260, 12.2, 4.58)."""
    if x is None or (isinstance(x, float) and math.isnan(x)):
        return ""
    if x == 0:
        return "0"
    d = max(0, 2 - int(math.floor(math.log10(abs(x)))))
    s = f"{x:.{d}f}"
    # rounding can add a digit (e.g. 9.996 -> 10.00); recompute once on the rounded value
    r = float(s)
    if r != 0:
        d2 = max(0, 2 - int(math.floor(math.log10(abs(r)))))
        s = f"{x:.{d2}f}" if d2 < d else s
    return s


def stars(p):
    return "***" if p < 0.01 else "**" if p < 0.05 else "*" if p < 0.10 else ""


def tstars(t):  # Table 2 high-minus-low tests: two-sided normal critical values
    a = abs(t)
    return "***" if a > 2.5758 else "**" if a > 1.96 else "*" if a > 1.645 else ""


def comma(n):
    return f"{int(round(n)):,}"


# ---------------------------------------------------------------- document helpers
doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(1.5))
base = doc.styles["Normal"]
base.font.name = "Times New Roman"
base.font.size = Pt(9)


def para(text, bold=False, size=9, italic=False, space_after=2):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def table(rows, font=7.5, first_col_width=None):
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, r in enumerate(rows):
        for j in range(ncol):
            cell = t.cell(i, j)
            cell.text = r[j] if j < len(r) else ""
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(font)
    if first_col_width:
        for row in t.rows:
            row.cells[0].width = Cm(first_col_width)
    return t


def caption_block(prefix, panel=None, notes=True):
    tx = text_for(prefix)
    para(tx["caption"], bold=True, size=10)
    if notes:
        for n in tx["notes"]:
            para(n, size=8, italic=True)
    if panel is not None:
        para(tx["panels"][panel], bold=True, size=9)


# ---------------------------------------------------------------- Table 1
t1 = pd.read_csv(out_dir / "results_table1.csv")
v1 = {(r.panel, str(r.row), r.stat): r.value for r in t1.itertuples()}
caption_block("Table 1 ", panel=0)
A = [("Skewness", "skew1"), ("Age", "age_ritter"), ("Coskewness", "coskewness"), ("Internet", "internet"),
     ("NASDAQ", "NASDAQ"), ("NYSE", "NYSE"), ("Ln(Proceeds)", "lnproceeds_adj1975"), ("Share Overhang", "share_overhang"),
     ("Pure Primary", "pure_primary"), ("Venture-Backed Deal", "VB"), ("Top-Tier Underwriter", "TT_ave"),
     ("Retail Investor Optimism", "lagmcsi"), ("Market Return", "lagcummktret"), ("R0", "r0"), ("R1", "r2_bhar_3y"),
     ("R2", "r3_bhar_3y")]
rows = [["Variables", "N", "Mean", "Standard deviation", "MAX", "MIN"]]
for lab, v in A:
    rows.append([lab, comma(v1[("A", v, "N")])] + [f"{v1[('A', v, s)]:.4f}" for s in ("Mean", "SD", "Max", "Min")])
table(rows, font=8)
para(text_for("Table 1 ")["panels"][1], bold=True, size=9)
FF30 = ["Food products", "Beer and liquor", "Tobacco products", "Recreation", "Printing and publishing", "Consumer goods",
        "Apparel", "Healthcare, medical equipment, pharmaceuticals", "Chemicals", "Textiles",
        "Construction and construction materials", "Steel works", "Fabricated products and machinery",
        "Electrical equipment", "Automobiles and trucks", "Aircraft, ships, and railroad equipment",
        "Precious metals, nonmetallic, and metal mining", "Coal", "Petroleum and natural gas", "Utilities",
        "Communication", "Personal and business services", "Business equipment",
        "Business supplies and shipping containers", "Transportation", "Wholesale", "Retail",
        "Restaurants, hotels, motels", "Banking, insurance, real estate, trading"]
rows = [["Industry", "", "Number of IPOs", "Mean expected skewness", "Low expected skewness",
         "Medium expected skewness", "High expected skewness"]]
for j, name in enumerate(FF30, start=1):
    k = str(j)
    rows.append([k, name, comma(v1[("B", k, "N")]), f"{v1[('B', k, 'Mean skew')]:.4f}"] +
                [f"{v1[('B', k, f'Tercile {g} %')]:.2f}%" for g in (1, 2, 3)])
table(rows, font=8)
doc.add_page_break()

# ---------------------------------------------------------------- Table 2
t2 = pd.read_csv(out_dir / "results_table2.csv")
v2 = {(str(r.row), r.stat): r.value for r in t2.itertuples()}
med = next(r for r in csv.DictReader(open(out_dir / "table2_median_params.csv")) if r["Parameter"] == "high_dummy")
med_t = float(med["tValue"])
caption_block("Table 2 ")
rows = [["Expected skewness", "Mean R0", "Median R0", "Observations", "Mean R1", "Mean R2"]]
for g, lab in (("1", "Low"), ("2", "Medium"), ("3", "High")):
    rows.append([lab, f"{v2[(g, 'Mean R0')]:.4f}", f"{v2[(g, 'Median R0')]:.4f}", comma(v2[(g, "Observations")]),
                 f"{v2[(g, 'Mean R1')]:.4f}", f"{v2[(g, 'Mean R2')]:.4f}"])
rows.append(["High-Low", f"{v2[('HL', 'diff r0')]:.4f}", f"{v2[('HL', 'diff median r0')]:.4f}", "",
             f"{v2[('HL', 'diff r2_bhar_3y')]:.4f}", f"{v2[('HL', 'diff r3_bhar_3y')]:.4f}"])
tt = lambda y: f"({v2[('HL', 't ' + y)]:.2f}){tstars(v2[('HL', 't ' + y)])}"
rows.append(["t-stat", tt("r0"), f"({med_t:.2f}){tstars(med_t)}", "", tt("r2_bhar_3y"), tt("r3_bhar_3y")])
table(rows, font=9)
doc.add_page_break()

# ---------------------------------------------------------------- regression tables
res = pd.concat([pd.read_csv(out_dir / f"results_{n}.csv") for n in ("returns", "panels", "revisions", "appendix_b", "ppml")])
res["panel"] = res["panel"].fillna(".").astype(str)
R = {}
for r in res.itertuples():
    m = R.setdefault((r.table, r.panel, int(r.column)), {"N": r.N, "fit": r.fit, "terms": {}})
    m["terms"][r.term] = (r.b, r.stat, r.p)

IPO = ["skew1", "age_ritter", "coskewness", "internet", "NASDAQ", "NYSE", "lnproceeds_adj1975", "price_adj_imp",
       "share_overhang", "pure_primary", "VB", "TT_ave", "lagmcsi", "lagcummktret", "ipo_vol", "ind_ret", "ind_mom",
       "ind_vol", "ind_turn", "_cons"]
REV = ["skew1", "age_ritter", "coskewness", "internet", "NASDAQ", "NYSE", "ln_amount_filed_adj1975", "share_overhang",
       "pure_primary", "VB", "uw_rank_ave", "lagmcsi", "lagcummktret", "ipo_vol", "ind_ret", "ind_mom", "ind_vol",
       "ind_turn", "_cons"]
PANEL = ["skew_l1", "skew_l2", "skew_l3", "skew_l4", "l_age", "l_cosk", "l_lnexpected", "l_overhang", "l_uwrank",
         "l_vcshare", "l_primaryshare", "l_optimism", "market_return_qminus1_candidate", "l_ipovol", "_cons"]
LAB = {"skew1": "Skewness", "age_ritter": "Age", "coskewness": "Coskewness", "internet": "Internet", "NASDAQ": "NASDAQ",
       "NYSE": "NYSE", "lnproceeds_adj1975": "Ln(Proceeds)", "price_adj_imp": "Price Adjustment",
       "share_overhang": "Share Overhang", "pure_primary": "Pure Primary", "VB": "Venture-Backed Deal",
       "TT_ave": "Top-Tier Underwriter", "lagmcsi": "Retail Investor Optimism", "lagcummktret": "Market Return",
       "ipo_vol": "IPO Volatility", "ind_ret": "Industry Return", "ind_mom": "Industry Momentum",
       "ind_vol": "Industry Volatility", "ind_turn": "Industry Turnover", "_cons": "Constant",
       "ln_amount_filed_adj1975": "Ln(Expected Proceeds)", "uw_rank_ave": "Underwriter Rank", "range_width": "Range Width",
       "skew_l1": "Skewness q−1", "skew_l2": "Skewness q−2", "skew_l3": "Skewness q−3", "skew_l4": "Skewness q−4",
       "l_age": "Age", "l_cosk": "Coskewness", "l_lnexpected": "Ln(Expected Proceeds)", "l_overhang": "Share Overhang",
       "l_uwrank": "Underwriter Rank", "l_vcshare": "VC-Backed Share", "l_primaryshare": "Pure Primary Share",
       "l_optimism": "Retail Investor Optimism", "market_return_qminus1_candidate": "Market Return",
       "l_ipovol": "IPO Volatility"}


def reg_table(cols, groups, terms, skew_label="Skewness", fit_label="Adjusted R-squared",
              fe_labels=("Year Fixed Effects", "Industry Fixed Effects"), year_from=4, ind_from=5, font=7):
    """cols: list of (table, panel, column). groups: [(header, n_cols)]."""
    head1 = [""] + [g for g, n in groups for _ in range(n)]
    head2 = ["VARIABLES"] + [f"({i})" for i in range(1, len(cols) + 1)]
    rows = [head1, head2]
    for term in terms:
        if not any(term in R[c]["terms"] for c in cols):
            continue
        lab = skew_label if term == "skew1" else LAB[term]
        b_row, t_row = [lab], [""]
        for c in cols:
            v = R[c]["terms"].get(term)
            b_row.append("" if v is None else sig3(v[0]) + stars(v[2]))
            t_row.append("" if v is None else f"({sig3(v[1])})")
        rows += [b_row, t_row]
    rows.append(["Observations"] + [comma(R[c]["N"]) for c in cols])
    rows.append([fit_label] + [f"{R[c]['fit']:.3f}" for c in cols])
    rows.append([fe_labels[0]] + ["YES" if c[2] >= year_from else "NO" for c in cols])
    rows.append([fe_labels[1]] + ["YES" if c[2] >= ind_from else "NO" for c in cols])
    table(rows, font=font, first_col_width=3.6)


def cols(table_, panel, n):
    return [(table_, panel, c) for c in range(1, n + 1)]


caption_block("Table 3 ")
reg_table(cols("Table3", "R0", 5) + cols("Table3", "R1", 5) + cols("Table3", "R2", 5), [("R0", 5), ("R1", 5), ("R2", 5)],
          IPO, font=6.5)
doc.add_page_break()
caption_block("Table 4 Expected", panel=0)
reg_table(cols("Table4", "A", 9), [("Ln(Primary Proceeds)", 9)], PANEL, fe_labels=("Year Effects", "Industry Fixed Effects"),
          year_from=5, ind_from=5)
doc.add_page_break()
caption_block("Table 4 continued", panel=0)
reg_table(cols("Table4", "B", 9), [("IPO Count", 9)], PANEL, fe_labels=("Year Effects", "Industry Fixed Effects"),
          year_from=5, ind_from=5)
doc.add_page_break()
caption_block("Table 5 Expected")
reg_table(cols("Table5", ".", 9), [("IPO Filing Count", 9)], PANEL, fe_labels=("Year Effects", "Industry Fixed Effects"),
          year_from=5, ind_from=5)
doc.add_page_break()
caption_block("Table 6 ")
reg_table(cols("Table6", ".", 5), [("Proceeds Revision", 5)], REV, fe_labels=("Year Effects", "Industry Fixed Effects"))
doc.add_page_break()
caption_block("Table 7 Expected", panel=0)
reg_table(cols("Table7", "A", 5), [("Share Revision", 5)], REV, fe_labels=("Year Effects", "Industry Fixed Effects"))
doc.add_page_break()
caption_block("Table 7 continued", panel=0)
reg_table(cols("Table7", "B", 5), [("Price Revision", 5)], REV, fe_labels=("Year Effects", "Industry Fixed Effects"))
doc.add_page_break()
caption_block("Table B1 ")
reg_table(cols("TableB1", "R1", 5) + cols("TableB1", "R2", 5), [("R1 1yr", 5), ("R2 1yr", 5)], IPO,
          skew_label="Expected Skewness")
doc.add_page_break()
caption_block("Table B2 ")
reg_table(cols("TableB2", "R1", 5) + cols("TableB2", "R2", 5), [("R1 5yr", 5), ("R2 5yr", 5)], IPO,
          skew_label="Expected Skewness")
doc.add_page_break()
caption_block("Table B3 ")
reg_table(cols("TableB3", ".", 5), [("Upward Price Revision", 5)], [t for t in REV if t != "_cons"],
          fit_label="Pseudo R-squared", fe_labels=("Year Effects", "Industry Fixed Effects"))
doc.add_page_break()
caption_block("Table B4 ")
reg_table(cols("TableB4", ".", 6), [("RPOS", 6)], ["skew1", "range_width"] + REV[1:],
          fe_labels=("Year Effects", "Industry Fixed Effects"), ind_from=5)
for prefix, key, head in (("Table IA1", "TableIA1", "Primary Proceeds"), ("Table IA2", "TableIA2", "IPO Count"),
                          ("Table IA3", "TableIA3", "IPO Filing Count")):
    doc.add_page_break()
    caption_block(prefix)
    reg_table(cols(key, ".", 9), [(head, 9)], PANEL, fit_label="Pseudo R-squared",
              fe_labels=("Year Effects", "Industry Fixed Effects"), year_from=5, ind_from=5)
doc.save(out_docx)
print("wrote", out_docx)
