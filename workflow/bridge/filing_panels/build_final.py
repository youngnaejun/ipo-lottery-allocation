from pathlib import Path
import pandas as pd, numpy as np, hashlib, json
from openpyxl import load_workbook
W=Path(__file__).resolve().parent
B=W.parent
ROOT=Path(r'@RUN_ROOT@')
X=ROOT/'stata_output/additional_tests/completion_industry_crosswalk/filing_completion_crosswalk_analysis.dta'
inputs=[X,B/'t11_yearfe_confirmed.dta',B/'t14_confirmed_panel.dta',B/'t11_three_definitions_deal_audit.csv',ROOT/'raw/@SDC_FULL_HISTORY_FILE@',ROOT/'raw/@SDC_ALL_FILINGS_FILE@']
def key(s):return s.astype(str).str.replace(r'\.0$','',regex=True)
def readxl(path):
 wb=load_workbook(path,read_only=True,data_only=True); it=wb.active.iter_rows(min_row=3,values_only=True); head=list(next(it));df=pd.DataFrame(it,columns=head);wb.close();return df.loc[:,~df.columns.duplicated()]
f=pd.read_stata(X,convert_categoricals=False);f['dk']=key(f.deal_key);assert (@PAPER_CHECKS@==0 or len(f)==17767) and f.dk.is_unique
raw=readxl(inputs[-1]);raw=raw[raw['SDC Deal Number'].notna()].copy();raw['dk']=key(raw['SDC Deal Number']);assert raw.dk.is_unique
allowed=set(f.security_type.unique())
trace=[]
def tr(stage,df):trace.append({'stage':stage,'n':len(df)})
r=raw;tr('raw pipeline extract',r)
for lab,mask in [('original IPO',r['Original IPO Flag'].eq('True')|r['Original IPO Flag'].eq(True))]:r=r[mask];tr(lab,r)
r=r[r['Issuer/Borrower Nation'].eq('United States')];tr('US issuer',r)
r=r[r['Security Type This Market'].isin(allowed)];tr('SAS common-equity allowlist',r)
r=r[r['Tranche Currency'].eq('US Dollar')];tr('USD',r)
assert set(r.dk)==set(f.dk),'SAS source population mismatch'
rawdates=pd.to_datetime(r['Dates: Filing Date'],format='%m/%d/%Y',errors='coerce')
rd=pd.Series(rawdates.to_numpy(),index=r.dk)
assert (f.dk.map(rd).to_numpy()==f.fdate.to_numpy()).all()
f['time_id']=(f.fdate.dt.year-1960)*4+f.fdate.dt.quarter-1;f['ff30']=f.ff30_conservative
f['is_spac']=f.spac_flag.astype(str).str.lower().eq('true')
d=pd.read_csv(inputs[3],dtype={'deal_num':str});f['in_t12']=f.dk.isin(d.deal_num)
a=f[f.fdate.dt.year.between(@FILING_START_YEAR@,@SAMPLE_END_YEAR@)].copy();tr('filing years @FILING_START_YEAR@-@SAMPLE_END_YEAR@',a)
a=a[a.ff30.between(1,29)];tr('FF30 1-29 mixed source',a)
a=a[~a.is_spac];tr('non-SPAC = prior A',a)
prior=a.copy();assert @PAPER_CHECKS@==0 or len(prior)==15977
# Immutable read-only reproduction comparison, before final population changes.
p=pd.read_stata(inputs[1],convert_categoricals=False)
p=p.merge(pd.read_stata(inputs[2],convert_categoricals=False)[['ff30','time_id','skew_l5']],on=['ff30','time_id'],validate='one_to_one')
assert (@PAPER_CHECKS@==0 or len(p)==4872) and not p.duplicated(['ff30','time_id']).any()
oldpath=B/'t14count/t14count_panel.dta';compared=False
if oldpath.exists():
 old=pd.read_stata(oldpath,convert_categoricals=False)
 for name,sel in {'nA':prior,'nA_comp':prior[prior.status_harmonised.eq('completed_live')],'nA_wd':prior[prior.status_harmonised.eq('withdrawn_cancelled')],'nB_comp':prior[prior.in_t12]}.items():
  cc=sel.groupby(['ff30','time_id']).size();check=old.set_index(['ff30','time_id']);assert np.array_equal(check[name],cc.reindex(check.index,fill_value=0))
 compared=True
# Evidence and screens use only source fields, never regression outcomes.
h=readxl(inputs[4]);h=h[h['SDC Deal Number'].notna()].copy();h['dk']=key(h['SDC Deal Number']);assert h.dk.is_unique
h['history_first']=pd.to_datetime(h['Full History Filing Date'].astype(str).str.split(';').str[0],format='%m/%d/%Y',errors='coerce')
f=f.merge(h[['dk','history_first','Full History Filing Date']],on='dk',how='left',validate='one_to_one')
f['history_agrees']=f.history_first.eq(f.fdate)
f['unknown_confirmed']=f.Transaction_Status.eq('Unknown')&f.history_agrees
f['status_allowed']=f.Transaction_Status.isin(['Live','Cancelled','Postponed','InProgress'])|f.unknown_confirmed
f['primary_industry']=f.SIC.fillna('').str.split(';').str[0].str.strip().str.lower()
f['reit']=f.SIC.fillna('').str.contains('real estate investment trusts',case=False,regex=False)|f.issuer_name.str.contains(r'\bREIT\b',case=False,regex=True)
fund_titles={'investment offices, nec','management investment offices, open-end','pension, health, and welfare funds','trusts, exc. educational, religious, & charitable','unit investment trusts, face-amount certificate offices'}
f['fund_proxy']=f.primary_industry.isin(fund_titles)
f['base_eligible']=f.fdate.dt.year.between(@FILING_START_YEAR@,@SAMPLE_END_YEAR@)&~f.is_spac
f['nonvehicle']=~f.reit&~f.fund_proxy
base=f[f.base_eligible&f.ff30.between(1,29)].copy()
a=base[base.status_allowed].copy();tr('actual-filing statuses; Unknown requires matching Full History date',a)
a=a[~a.reit];tr('exclude explicit REIT',a)
a=a[~a.fund_proxy];tr('exclude pooled-investment industry proxies = MAIN',a)
a=a.sort_values(['issuer_key','fdate','dk']);a['first_in_main']=~a.duplicated('issuer_key')
# This is a name-key robustness, not proof that every repeat is a refiling after withdrawal.
a['same_day_repeat']=a.duplicated(['issuer_key','fdate'],keep=False)
a['prior_status']=a.groupby('issuer_key').Transaction_Status.shift()
a['after_cancelled']=a.prior_status.eq('Cancelled')
statuses={'comp':'Live','wd':'Cancelled','pp':'Postponed','open':'InProgress','unk':'Unknown'}
counts={}
def add(name,df,col='ff30'):
 counts[name]=df.groupby([col,'time_id']).size().rename_axis(['ff30','time_id'])
add('n_main',a)
for k,v in statuses.items():add('n_'+k,a[a.Transaction_Status.eq(v)])
add('n_no_unk',a[a.Transaction_Status.ne('Unknown')])
add('n_all_unk',base[base.nonvehicle&base.Transaction_Status.isin(['Live','Cancelled','Postponed','InProgress','Unknown'])])
add('n_no_fund_screen',base[base.status_allowed&~base.reit])
add('n_priorA',prior)
add('n_bridge',a[a.in_t12|a.Transaction_Status.ne('Live')])
add('n_day',a.drop_duplicates(['issuer_key','fdate']));add('n_first',a[a.first_in_main]);add('n_first_wd',a[a.first_in_main&a.Transaction_Status.eq('Cancelled')])
common=a[a.fallback_ff30_conservative.between(1,29)].copy()
add('n_common_mix',common);add('n_common_sdc',common,'fallback_ff30_conservative')
add('n_cmix_wd',common[common.Transaction_Status.eq('Cancelled')]);add('n_csdc_wd',common[common.Transaction_Status.eq('Cancelled')],'fallback_ff30_conservative')
u=f[f.base_eligible&f.nonvehicle&f.status_allowed&f.fallback_ff30_conservative.between(1,29)]
add('n_uniform',u,'fallback_ff30_conservative');add('n_uniform_wd',u[u.Transaction_Status.eq('Cancelled')],'fallback_ff30_conservative')
idx=pd.MultiIndex.from_frame(p[['ff30','time_id']])
for n,s in counts.items():
 assert s.index.isin(idx).all()
 p[n]=s.reindex(idx,fill_value=0).to_numpy().astype(int)
assert np.array_equal(p.n_main,p[['n_comp','n_wd','n_pp','n_open','n_unk']].sum(axis=1))
# Controls and skewness unaltered; T12 remains exact in the preserved bridge audit.
p.to_stata(W/'T14_Filing_Panel_20260924.dta',write_index=False,version=118)
keep=['dk','issuer_name','issuer_key','fdate','Transaction_Status','status_harmonised','ff30','fallback_ff30_conservative','industry_source_conservative','confidence','match_method','history_first','history_agrees','unknown_confirmed','reit','fund_proxy','primary_industry','in_t12']
f[keep].to_csv(W/'T14_Source_Deal_Audit.csv',index=False)
a[keep+['first_in_main','same_day_repeat','prior_status','after_cancelled']].to_csv(W/'T14_Final_Filings.csv',index=False)
base[base.Transaction_Status.eq('Unknown')][keep].to_csv(W/'T14_Unknown_Audit.csv',index=False)
pd.DataFrame(trace).to_csv(W/'T14_Screen_Trace.csv',index=False)
a.groupby([a.fdate.dt.year,'Transaction_Status']).size().unstack(fill_value=0).to_csv(W/'T14_Status_By_Year.csv')
pd.crosstab(a.Transaction_Status,a.industry_source_conservative).to_csv(W/'T14_Industry_Source.csv')
ex=f[f.base_eligible&f.ff30.between(1,29)].copy();ex['in_main']=ex.dk.isin(a.dk);ex.loc[~ex.in_main,keep].to_csv(W/'T14_Excluded_Deals.csv',index=False)
audit={'prior_build_compared':compared,'prior_build_passed':compared,'raw_date_matches_master':True,'counts':{n:int(p[n].sum()) for n in counts},'statuses':a.Transaction_Status.value_counts().to_dict(),'common_deals':len(common),'common_mapping_agreement':float(common.ff30.eq(common.fallback_ff30_conservative).mean()),'unknown_in_priorA':int(base.Transaction_Status.eq('Unknown').sum()),'unknown_confirmed_priorA':int(base.unknown_confirmed.sum()),'nonfiling_priorA':int(base.Transaction_Status.isin(['Rumoured','Announced','Mandated']).sum()),'repeated_names_main':int((~a.first_in_main).sum()),'same_day_repeat_rows':int(a.same_day_repeat.sum()),'after_cancelled_name_key':int(a.after_cancelled.sum()),'main_T12_overlap':int(a.in_t12.sum()),'source_status_date':'2026-08-12','history_date':'2026-09-03','grid_cells':len(p)}
(W/'T14_Audit.json').write_text(json.dumps(audit,indent=2),encoding='utf-8')
manifest=[{'path':str(v),'sha256':hashlib.sha256(v.read_bytes()).hexdigest(),'bytes':v.stat().st_size} for v in inputs]
(W/'T14_Input_Manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
print(json.dumps(audit,indent=2));print(pd.DataFrame(trace).to_string(index=False))

