from pathlib import Path
import pandas as pd,numpy as np,json
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[1]
O=R/'results/all_filer_controls';O.mkdir(parents=True,exist_ok=True)
S=Path('@RUN_ROOT@/ext/ipo-20260921-25/t14filing_final')
P=Path('@RUN_ROOT@')
f=pd.read_csv(S/'T14_Final_Filings.csv',dtype={'dk':str})
f['fdate']=pd.to_datetime(f.fdate);f['time_id']=(f.fdate.dt.year-1960)*4+f.fdate.dt.quarter-1
assert (@PAPER_CHECKS@==0 or len(f)==13957) and f.dk.is_unique
w=load_workbook(P/'raw/@SDC_FULL_HISTORY_FILE@',read_only=True,data_only=True)
it=w.active.iter_rows(min_row=3,values_only=True);head=list(next(it))
names=['SDC Deal Number','Full History Filing Date','Full History Amount Filed This Market (USD Millions)','Full History Total Shares Filed This Market','Full History Primary Shares Filed This Market','Venture Capital Backed IPO Issue Flag','Lead Managers','Lead Managers Code']
idx=[head.index(n) for n in names];h=pd.DataFrame(([r[i] for i in idx] for r in it),columns=names);w.close()
h=h[h[names[0]].notna()].copy();h['dk']=h[names[0]].astype(str).str.replace(r'\.0$','',regex=True)
assert h.dk.is_unique
d=f.merge(h,on='dk',how='left',validate='one_to_one')
def first(x):
 if pd.isna(x):return np.nan
 return str(x).split(';')[0].strip()
def numeric(x):
 v=first(x)
 return pd.to_numeric(str(v).replace(',','').replace('$',''),errors='coerce')
d['first_history_date']=pd.to_datetime(d[names[1]].map(first),format='%m/%d/%Y',errors='coerce')
d['history_aligned']=d.first_history_date.eq(d.fdate)
for name,col in [('initial_amount',names[2]),('initial_total',names[3]),('initial_primary',names[4])]:d[name]=d[col].map(numeric).where(d.history_aligned)
d['ln_initial_amount']=np.log(d.initial_amount.where(d.initial_amount>0))
# Zero first-history primary quantities are ambiguous in the existing source audit.
valid=(d.initial_total>0)&(d.initial_primary>0)&(d.initial_primary<=d.initial_total)
d['initial_pure_primary']=d.initial_primary.eq(d.initial_total).astype(float).where(valid)
d['vc_snapshot']=d[names[5]].astype(str).str.lower().map({'true':1.,'false':0.,'yes':1.,'no':0.})
d.to_csv(O/'Deal_Control_Audit.csv',index=False)
cov=[]
for status,g in [('ALL',d)]+list(d.groupby('Transaction_Status')):
 for v in ['history_aligned','ln_initial_amount','initial_pure_primary','vc_snapshot']:
  n=int(g[v].sum()) if v=='history_aligned' else int(g[v].notna().sum())
  cov.append(dict(status=status,variable=v,n=len(g),available=n,coverage=n/len(g)))
c=pd.DataFrame(cov);c.to_csv(O/'Coverage_By_Status.csv',index=False)
print(c.to_string(index=False))
print('VC raw categories',d[names[5]].value_counts(dropna=False).to_dict())
